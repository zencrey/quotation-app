import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "data" / "quotations.xlsx"

# Header styling
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
REGULAR_FONT = Font(name="Calibri", size=11)
BORDER_THIN = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

QUOTATION_HEADERS = [
    "Quotation ID",
    "Quotation Number",
    "Date",
    "Customer Name",
    "Total Amount",
    "Amount in Words",
    "Incharge",
    "Email",
    "Phone Number"
]

ITEM_HEADERS = [
    "Quotation ID",
    "Description",
    "Quantity",
    "Unit Price",
    "Line Total"
]

def init_excel_file():
    """Initializes the Excel workbook if it does not already exist."""
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    if not EXCEL_PATH.exists():
        wb = openpyxl.Workbook()
        
        # Setup Quotations sheet
        ws_quotations = wb.active
        ws_quotations.title = "Quotations"
        ws_quotations.append(QUOTATION_HEADERS)
        for col_idx in range(1, len(QUOTATION_HEADERS) + 1):
            cell = ws_quotations.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_THIN
        
        # Setup Quotation Items sheet
        ws_items = wb.create_sheet(title="Quotation Items")
        ws_items.append(ITEM_HEADERS)
        for col_idx in range(1, len(ITEM_HEADERS) + 1):
            cell = ws_items.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_THIN
            
        auto_fit_columns(ws_quotations)
        auto_fit_columns(ws_items)
        wb.save(EXCEL_PATH)

def auto_fit_columns(ws):
    """Adjusts column widths based on content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

def save_quotation_to_excel(quotation, items):
    """
    Saves or updates a quotation and its items in data/quotations.xlsx.
    Preserves all other quotations in the workbook.
    """
    init_excel_file()
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # 1. Update or append in 'Quotations' worksheet
    if "Quotations" not in wb.sheetnames:
        ws_quotations = wb.create_sheet(title="Quotations")
        ws_quotations.append(QUOTATION_HEADERS)
    else:
        ws_quotations = wb["Quotations"]

    target_q_row = None
    for row in range(2, ws_quotations.max_row + 1):
        cell_id = ws_quotations.cell(row=row, column=1).value
        if cell_id is not None and str(cell_id) == str(quotation.id):
            target_q_row = row
            break

    q_data = [
        quotation.id,
        quotation.quotation_no,
        quotation.date,
        quotation.customer_name,
        quotation.total_amount,
        quotation.amount_in_words,
        quotation.incharge or "",
        quotation.email or "",
        quotation.phone or ""
    ]

    if target_q_row:
        # Overwrite existing row
        for col_idx, val in enumerate(q_data, start=1):
            cell = ws_quotations.cell(row=target_q_row, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN
    else:
        # Append new row
        ws_quotations.append(q_data)
        new_row = ws_quotations.max_row
        for col_idx in range(1, len(q_data) + 1):
            cell = ws_quotations.cell(row=new_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN

    # 2. Update 'Quotation Items' worksheet
    if "Quotation Items" not in wb.sheetnames:
        ws_items = wb.create_sheet(title="Quotation Items")
        ws_items.append(ITEM_HEADERS)
    else:
        ws_items = wb["Quotation Items"]

    # Remove existing items for this quotation_id (scan backwards)
    for row in range(ws_items.max_row, 1, -1):
        cell_id = ws_items.cell(row=row, column=1).value
        if cell_id is not None and str(cell_id) == str(quotation.id):
            ws_items.delete_rows(row)

    # Append new items
    for item in items:
        unit_price = item.unit_price_kd + (item.unit_price_fills / 1000.0)
        line_total = item.total_price if item.total_price else (item.total_kd + item.total_fills / 1000.0)
        
        item_data = [
            quotation.id,
            item.description,
            item.quantity,
            unit_price,
            line_total
        ]
        ws_items.append(item_data)
        item_row = ws_items.max_row
        for col_idx in range(1, len(item_data) + 1):
            cell = ws_items.cell(row=item_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN

    auto_fit_columns(ws_quotations)
    auto_fit_columns(ws_items)
    wb.save(EXCEL_PATH)

def delete_quotation_from_excel(quotation_id):
    """Removes a quotation and its items from data/quotations.xlsx."""
    if not EXCEL_PATH.exists():
        return
    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    if "Quotations" in wb.sheetnames:
        ws_quotations = wb["Quotations"]
        for row in range(ws_quotations.max_row, 1, -1):
            if str(ws_quotations.cell(row=row, column=1).value) == str(quotation_id):
                ws_quotations.delete_rows(row)
                
    if "Quotation Items" in wb.sheetnames:
        ws_items = wb["Quotation Items"]
        for row in range(ws_items.max_row, 1, -1):
            if str(ws_items.cell(row=row, column=1).value) == str(quotation_id):
                ws_items.delete_rows(row)
                
    wb.save(EXCEL_PATH)
